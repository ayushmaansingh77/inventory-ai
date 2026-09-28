import csv
import io

from openpyxl import load_workbook

from app import db
from app.models.inventory import InventoryItem

ALLOWED_EXTENSIONS = {"csv", "xlsx"}

REQUIRED_FIELDS = ["name", "sku", "quantity", "unit_price"]
OPTIONAL_FIELDS = ["reorder_level"]


def _get_extension(filename):
    if not filename or "." not in filename:
        return None
    return filename.rsplit(".", 1)[1].lower()


def parse_spreadsheet(file_storage):
    """
    Reads an uploaded CSV or XLSX file into (headers, rows, error).

    `rows` is a list of dicts keyed by the raw spreadsheet header strings,
    with every value as a plain string — field-name mapping and type
    coercion happen later in import_inventory_items, once the caller has
    told us which header maps to which inventory field.
    """
    ext = _get_extension(file_storage.filename)

    if ext not in ALLOWED_EXTENSIONS:
        return None, None, "File must be a .csv or .xlsx spreadsheet"

    try:
        if ext == "csv":
            text = file_storage.read().decode("utf-8-sig")
            rows_raw = list(csv.reader(io.StringIO(text)))
        else:
            workbook = load_workbook(file_storage, read_only=True, data_only=True)
            sheet = workbook.active
            rows_raw = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
    except Exception:
        return None, None, "Could not read the file — make sure it's a valid CSV or XLSX"

    if len(rows_raw) < 2:
        return None, None, "File must contain a header row and at least one data row"

    headers = [str(h).strip() for h in rows_raw[0]]

    data_rows = []
    for raw_row in rows_raw[1:]:
        if all(str(cell).strip() == "" for cell in raw_row):
            continue  # skip fully blank rows

        row_dict = {}
        for i, header in enumerate(headers):
            row_dict[header] = str(raw_row[i]).strip() if i < len(raw_row) else ""
        data_rows.append(row_dict)

    return headers, data_rows, None


def import_inventory_items(user_id, rows, column_mapping):
    """
    rows: list of dicts keyed by raw spreadsheet header, as returned by
        parse_spreadsheet.
    column_mapping: dict mapping our inventory field names to the
        spreadsheet header that holds that data, e.g.
        {"name": "Product Name", "sku": "SKU Code", "quantity": "Qty",
         "unit_price": "Price", "reorder_level": "Reorder Point"}
        (reorder_level is optional; if omitted or blank, defaults to 10,
        matching the existing single-item create default.)

    A row whose SKU already exists for this user (in the DB or earlier
    in this same file) is skipped, not overwritten, and reported back so
    nothing is silently lost or clobbered.

    Returns (summary, error). summary = {
        "created": int,
        "skipped_duplicate_skus": [sku, ...],
        "row_errors": [{"row": row_number, "error": "..."}],
    }
    """
    for field in REQUIRED_FIELDS:
        if not column_mapping.get(field):
            return None, f"Column mapping is missing a required field: {field}"

    existing_skus = {
        sku
        for (sku,) in db.session.query(InventoryItem.sku)
        .filter_by(user_id=user_id)
        .all()
    }

    created = 0
    skipped_duplicate_skus = []
    row_errors = []
    new_items = []

    for i, row in enumerate(rows):
        row_number = i + 2  # +1 for the header row, +1 to 1-index

        name = row.get(column_mapping["name"], "").strip()
        sku = row.get(column_mapping["sku"], "").strip()
        quantity_raw = row.get(column_mapping["quantity"], "").strip()
        unit_price_raw = row.get(column_mapping["unit_price"], "").strip()
        reorder_header = column_mapping.get("reorder_level")
        reorder_raw = row.get(reorder_header, "").strip() if reorder_header else ""

        if not name or not sku:
            row_errors.append({"row": row_number, "error": "Missing name or SKU"})
            continue

        if sku in existing_skus:
            skipped_duplicate_skus.append(sku)
            continue

        try:
            quantity = int(float(quantity_raw)) if quantity_raw else 0
            unit_price = float(unit_price_raw)
            reorder_level = int(float(reorder_raw)) if reorder_raw else 10
        except ValueError:
            row_errors.append({
                "row": row_number,
                "error": "Quantity, unit price, or reorder level is not a valid number",
            })
            continue

        new_items.append(
            InventoryItem(
                name=name,
                sku=sku,
                quantity=quantity,
                unit_price=unit_price,
                reorder_level=reorder_level,
                user_id=user_id,
            )
        )
        existing_skus.add(sku)  # guards against duplicate SKUs within this same file
        created += 1

    if new_items:
        db.session.add_all(new_items)
        db.session.commit()

    return {
        "created": created,
        "skipped_duplicate_skus": skipped_duplicate_skus,
        "row_errors": row_errors,
    }, None
