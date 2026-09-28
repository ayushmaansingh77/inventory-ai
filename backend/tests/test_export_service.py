import csv
import io

from openpyxl import load_workbook

from app.services.export_service import generate_inventory_export


def test_rejects_unsupported_format(app, user):
    content, mimetype, filename, error = generate_inventory_export(user.id, "pdf")
    assert content is None
    assert "csv" in error and "xlsx" in error


def test_csv_export_contains_header_and_items(app, user, item_without_history):
    content, mimetype, filename, error = generate_inventory_export(user.id, "csv")

    assert error is None
    assert mimetype == "text/csv"
    assert filename == "inventory_export.csv"

    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert rows[0] == ["Name", "SKU", "Quantity", "Unit Price", "Reorder Level"]
    assert rows[1][1] == "TEST-002"  # item_without_history's sku


def test_csv_export_empty_inventory_has_only_header(app, user):
    content, mimetype, filename, error = generate_inventory_export(user.id, "csv")

    assert error is None
    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert len(rows) == 1


def test_xlsx_export_contains_header_and_items(app, user, item_without_history):
    content, mimetype, filename, error = generate_inventory_export(user.id, "xlsx")

    assert error is None
    assert mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert filename == "inventory_export.xlsx"

    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Name", "SKU", "Quantity", "Unit Price", "Reorder Level")
    assert rows[1][1] == "TEST-002"


def test_export_only_includes_current_users_items(app, user, item_without_history):
    from app import db
    from app.models.user import User
    from app.models.inventory import InventoryItem

    other_user = User(username="other", email="other@example.com", is_verified=True)
    other_user.set_password("password123")
    db.session.add(other_user)
    db.session.commit()

    other_item = InventoryItem(
        name="Other User's Item", sku="OTHER-1", quantity=1, unit_price=1.0, user_id=other_user.id
    )
    db.session.add(other_item)
    db.session.commit()

    content, mimetype, filename, error = generate_inventory_export(user.id, "csv")
    assert error is None
    assert "OTHER-1" not in content.decode("utf-8")
